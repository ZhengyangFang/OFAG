"""Rows of named values, from whatever file they arrived in."""

import csv
from dataclasses import dataclass
from importlib.util import find_spec
from io import StringIO
from pathlib import Path

#: Tried in order when the caller does not state one.
TEXT_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "cp1252")

#: Extension to the delimiter that extension promises.
_DELIMITERS: dict[str, str] = {".csv": ",", ".txt": ",", ".tsv": "\t", ".tab": "\t"}

_EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm"})
_VECTOR_SUFFIXES = frozenset({".dbf", ".shp", ".gpkg", ".geojson", ".json"})


@dataclass(frozen=True)
class Table:
    """A header, its rows, and how they were read."""

    columns: tuple[str, ...]
    rows: tuple[dict[str, str], ...]
    #: Which branch read it, for the record a caller keeps of its inputs.
    source_format: str
    #: The text encoding used, or None for a format that is not text.
    encoding: str | None
    #: The coordinate system the file declares, for the formats that carry one.
    crs: str | None = None

    def __len__(self) -> int:
        return len(self.rows)


def supported_suffixes() -> tuple[str, ...]:
    """Every extension this can read here and now, dependencies included."""
    suffixes = set(_DELIMITERS)
    if find_spec("openpyxl") is not None:
        suffixes |= _EXCEL_SUFFIXES
    if find_spec("pyogrio") is not None:
        suffixes |= _VECTOR_SUFFIXES
    return tuple(sorted(suffixes))


def read_table(path: Path | str, *, encoding: str | None = None) -> Table:
    """One file as rows, dispatched on its extension."""
    source = Path(path)
    if not source.is_file():
        raise ValueError(f"file does not exist: {source}")
    suffix = source.suffix.lower()
    if suffix in _DELIMITERS:
        return _read_delimited(source, _DELIMITERS[suffix], encoding)
    if suffix in _EXCEL_SUFFIXES:
        return _read_excel(source)
    if suffix in _VECTOR_SUFFIXES:
        return _read_vector(source)
    known = ", ".join(supported_suffixes())
    raise ValueError(f"cannot read {suffix or 'a file with no extension'}; readable here: {known}")


def read_text_table(text: str, *, delimiter: str = ",") -> Table:
    """Rows from text already in hand, for data that arrived in a request."""
    reader = csv.DictReader(StringIO(text), delimiter=delimiter)
    return Table(
        columns=tuple(reader.fieldnames or ()),
        rows=tuple(dict(row) for row in reader),
        source_format="text",
        encoding=None,
    )


def _read_delimited(path: Path, delimiter: str, encoding: str | None) -> Table:
    """Separated text, with the encoding detected when it is not stated."""
    candidates = (encoding,) if encoding is not None else TEXT_ENCODINGS
    last: UnicodeDecodeError | None = None
    for candidate in candidates:
        try:
            with path.open("r", encoding=candidate, newline="") as stream:
                reader = csv.DictReader(stream, delimiter=delimiter)
                columns = tuple(reader.fieldnames or ())
                rows = tuple(dict(row) for row in reader)
        except UnicodeDecodeError as error:
            last = error
            continue
        return Table(columns=columns, rows=rows, source_format="delimited", encoding=str(candidate))
    raise ValueError(f"{path} could not be decoded as {', '.join(map(str, candidates))}: {last}")


def _read_excel(path: Path) -> Table:
    """The first worksheet, with its first row as the header."""
    if find_spec("openpyxl") is None:
        raise ValueError(
            f"reading {path.suffix} needs openpyxl; install the optional extra with "
            "`uv sync --extra tables`"
        )
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    # read_only so a large export is streamed, data_only so a formula comes back as the
    # number it evaluated to rather than its source text.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        values = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not values:
        return Table(columns=(), rows=(), source_format="excel", encoding=None)
    columns = tuple("" if cell is None else str(cell).strip() for cell in values[0])
    rows = tuple(
        {name: "" if cell is None else str(cell) for name, cell in zip(columns, row, strict=False)}
        for row in values[1:]
    )
    return Table(columns=columns, rows=rows, source_format="excel", encoding=None)


def _read_vector(path: Path) -> Table:
    """A vector dataset's attribute table, geometry left behind."""
    if find_spec("pyogrio") is None:
        raise ValueError(
            f"reading {path.suffix} needs pyogrio; install the optional extra with "
            "`uv sync --extra tables`"
        )

    # The raw reader rather than `read_dataframe`, which needs geopandas for a frame
    # this immediately flattens back into rows.
    from pyogrio.raw import read  # type: ignore[import-untyped]

    meta, _, _, field_data = read(path, read_geometry=False)
    columns = tuple(str(name) for name in meta["fields"])
    rows = tuple(
        {
            name: "" if values[index] is None else str(values[index])
            for name, values in zip(columns, field_data, strict=True)
        }
        for index in range(len(field_data[0]) if field_data else 0)
    )
    return Table(
        columns=columns,
        rows=rows,
        source_format="vector",
        encoding=meta.get("encoding"),
        crs=meta.get("crs"),
    )
